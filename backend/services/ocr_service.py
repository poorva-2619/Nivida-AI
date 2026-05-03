import fitz  # PyMuPDF
import pytesseract
from PIL import Image
import docx
import io
import re

class OCRService:
    @staticmethod
    def extract_text_from_file(file_path: str, file_type: str) -> dict:
        result = {
            "text": "",
            "confidence": 0.0,
            "method": "",
            "pages": 0,
            "page_texts": [],
            "warnings": [],
            "language_detected": "mixed"
        }
        
        file_type = file_type.lower()
        
        if file_type == 'pdf':
            doc = fitz.open(file_path)
            result["pages"] = len(doc)
            
            full_text = ""
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text()
                result["page_texts"].append(text)
                full_text += text + "\n"
                
            if len(full_text.strip()) > 100:
                result["text"] = full_text
                result["confidence"] = 0.99
                result["method"] = "pdf_direct"
            else:
                result["method"] = "ocr_tesseract"
                result["page_texts"] = []
                full_text = ""
                total_conf = 0
                valid_pages = 0
                
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    pix = page.get_pixmap()
                    img_bytes = pix.tobytes("png")
                    img = Image.open(io.BytesIO(img_bytes))
                    
                    try:
                        text = pytesseract.image_to_string(img, lang="eng+hin+kan")
                        result["page_texts"].append(text)
                        full_text += text + "\n"
                        
                        data = pytesseract.image_to_data(img, lang="eng+hin+kan", output_type=pytesseract.Output.DICT)
                        confidences = [int(c) for c in data['conf'] if str(c).isdigit() and int(c) != -1]
                        
                        if confidences:
                            page_conf = sum(confidences) / len(confidences) / 100.0
                            total_conf += page_conf
                            valid_pages += 1
                            
                            if page_conf < 0.50:
                                result["warnings"].append(f"Low confidence on page {page_num + 1} - possible blurry scan")
                    except Exception as e:
                        result["warnings"].append(f"OCR failed on page {page_num + 1}: {str(e)}")
                
                result["text"] = full_text
                if valid_pages > 0:
                    result["confidence"] = total_conf / valid_pages
                else:
                    result["confidence"] = 0.0

        elif file_type in ['jpg', 'jpeg', 'png']:
            result["method"] = "ocr_tesseract"
            result["pages"] = 1
            try:
                img = Image.open(file_path)
                text = pytesseract.image_to_string(img, lang="eng+hin+kan")
                result["text"] = text
                result["page_texts"] = [text]
                
                data = pytesseract.image_to_data(img, lang="eng+hin+kan", output_type=pytesseract.Output.DICT)
                confidences = [int(c) for c in data['conf'] if str(c).isdigit() and int(c) != -1]
                
                if confidences:
                    result["confidence"] = sum(confidences) / len(confidences) / 100.0
                else:
                    result["confidence"] = 0.0
            except Exception as e:
                result["warnings"].append(f"OCR failed: {str(e)}")
                
        elif file_type == 'docx':
            result["method"] = "docx_extract"
            try:
                doc = docx.Document(file_path)
                full_text = []
                for para in doc.paragraphs:
                    full_text.append(para.text)
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            full_text.append(cell.text)
                            
                combined_text = "\n".join(full_text)
                result["text"] = combined_text
                result["page_texts"] = [combined_text]
                result["pages"] = 1
                result["confidence"] = 0.99
            except Exception as e:
                result["warnings"].append(f"DOCX extraction failed: {str(e)}")
                
        conf = result["confidence"]
        if conf < 0.50 and result["method"] != "":
            result["warnings"].append(f"Overall confidence is low ({conf:.2f}). Flagged for human review.")
        elif 0.50 <= conf <= 0.84:
            result["warnings"].append(f"Overall confidence is medium ({conf:.2f}). Proceed with caution.")
            
        return result

    @staticmethod
    def highlight_relevant_values(text: str, criteria_list: list) -> dict:
        highlighted_sections = []
        
        for criteria in criteria_list:
            criteria_id = criteria.get("id", "C_UNKNOWN")
            title = criteria.get("title", "").lower()
            
            matches = []
            if "turnover" in title or "financial" in title:
                pattern = r'(Rs\.?|INR|\$)\s*[\d,]+(\.\d+)?\s*(Crore|Cr|Lakh|Million|M)?'
                for m in re.finditer(pattern, text, re.IGNORECASE):
                    matches.append(m.group(0))
            elif "date" in title:
                pattern = r'\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b'
                for m in re.finditer(pattern, text):
                    matches.append(m.group(0))
            elif "certificate" in title or "iso" in title:
                pattern = r'(ISO|Certificate)\s*[:#-]?\s*[A-Z0-9-]{5,}'
                for m in re.finditer(pattern, text, re.IGNORECASE):
                    matches.append(m.group(0))
            
            for match in matches:
                highlighted_sections.append({
                    "criteria_id": criteria_id,
                    "matched_text": match,
                    "page": 1,
                    "confidence": 0.85
                })
                
        return {"highlighted_sections": highlighted_sections}
