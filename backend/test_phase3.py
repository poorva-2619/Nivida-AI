import requests
import json
import os
import docx

BASE_URL = "http://127.0.0.1:8000"

print("==========================================")
print("   Phase 3: OCR Pipeline Demonstration    ")
print("==========================================")

# 1. Login as Admin
print("\n[1] Authenticating as Admin...")
login_data = {
    "username": "admin@nividaai.gov.in",
    "password": "admin123"
}
response = requests.post(f"{BASE_URL}/api/auth/login", data=login_data)
if response.status_code != 200:
    print("Login failed!", response.text)
    exit(1)

token = response.json().get("access_token")
headers = {"Authorization": f"Bearer {token}"}
print(" -> Success! Acquired JWT Bearer Token.")

# 2. Create a dummy docx file
print("\n[2] Generating a sample Tender Document (DOCX)...")
doc = docx.Document()
doc.add_heading('NividaAI Test Tender Document', 0)
doc.add_paragraph("This is a sample tender document generated to test the OCR pipeline.")
doc.add_paragraph("Requirement 1: Minimum Annual Turnover must be Rs 10 Crore.")
doc.add_paragraph("Requirement 2: Bidding vendor must possess ISO-9001 Certificate.")
dummy_path = "test_tender.docx"
doc.save(dummy_path)
print(f" -> Created file: {dummy_path}")

# 3. Upload Tender
print("\n[3] Triggering POST /api/tender/upload to test ingestion and OCR...")
with open(dummy_path, "rb") as f:
    files = {"file": ("test_tender.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {"title": "Sample Road Construction Tender 2026"}
    upload_res = requests.post(f"{BASE_URL}/api/tender/upload", headers=headers, files=files, data=data)

print(f"\nSERVER RESPONSE (Status {upload_res.status_code}):")
print(json.dumps(upload_res.json(), indent=2))

# Cleanup
if os.path.exists(dummy_path):
    os.remove(dummy_path)
print("\n==========================================")
print("           Demonstration Complete           ")
print("==========================================")
